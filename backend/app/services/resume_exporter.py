"""Resume Exporter - Generates PDF and DOCX from resume data."""

from io import BytesIO
from app.services.resume_builder import _get_mock_resume


async def export_resume_pdf(resume_id: str) -> bytes:
    resume = _get_mock_resume("job-1", "modern")
    html = _build_html(resume)
    try:
        from weasyprint import HTML
        return HTML(string=html).write_pdf()
    except ImportError:
        return _simple_text(resume).encode("utf-8")


async def export_resume_docx(resume_id: str) -> bytes:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

    resume = _get_mock_resume("job-1", "modern")
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    heading = doc.add_heading(level=0)
    heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = heading.add_run("Rahul Kumar")
    run.font.size = Pt(24)

    contact = doc.add_paragraph()
    contact.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    contact.add_run("rahul@email.com | +91-9876543210 | Bangalore, India")
    doc.add_paragraph()

    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(resume.sections.summary)

    doc.add_heading("Experience", level=1)
    for exp in resume.sections.experience:
        p = doc.add_paragraph()
        run = p.add_run(exp.role)
        run.bold = True
        p.add_run(f" - {exp.company} ({exp.duration})")
        for bullet in exp.bullets:
            doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("Skills", level=1)
    doc.add_paragraph(", ".join(resume.sections.skills))

    doc.add_heading("Education", level=1)
    for edu in resume.sections.education:
        p = doc.add_paragraph()
        run = p.add_run(edu.degree)
        run.bold = True
        p.add_run(f" - {edu.college} ({edu.year})")

    if resume.sections.certifications:
        doc.add_heading("Certifications", level=1)
        for cert in resume.sections.certifications:
            doc.add_paragraph(cert, style="List Bullet")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def _build_html(resume) -> str:
    exp_html = ""
    for exp in resume.sections.experience:
        bullets = "".join(f"<li>{b}</li>" for b in exp.bullets)
        exp_html += f'<div><h3>{exp.role} <span style="color:#666">- {exp.company}</span></h3><p style="color:#888">{exp.duration}</p><ul>{bullets}</ul></div>'

    return f"""<!DOCTYPE html><html><head><style>
body{{font-family:Calibri,sans-serif;margin:40px;line-height:1.5;color:#333}}
h1{{text-align:center;margin-bottom:5px}}
.contact{{text-align:center;color:#666;margin-bottom:20px}}
h2{{color:#2563eb;border-bottom:2px solid #2563eb;padding-bottom:5px}}
</style></head><body>
<h1>Rahul Kumar</h1>
<p class="contact">rahul@email.com | +91-9876543210 | Bangalore, India</p>
<h2>Professional Summary</h2><p>{resume.sections.summary}</p>
<h2>Experience</h2>{exp_html}
<h2>Skills</h2><p>{", ".join(resume.sections.skills)}</p>
<h2>Education</h2><p><strong>{resume.sections.education[0].degree}</strong> - {resume.sections.education[0].college} ({resume.sections.education[0].year})</p>
<h2>Certifications</h2><ul>{"".join(f"<li>{c}</li>" for c in resume.sections.certifications)}</ul>
</body></html>"""


def _simple_text(resume) -> str:
    content = f"RAHUL KUMAR\nrahul@email.com | +91-9876543210\n\nSUMMARY\n{resume.sections.summary}\n\nEXPERIENCE\n"
    for exp in resume.sections.experience:
        content += f"\n{exp.role} - {exp.company} ({exp.duration})\n"
        for b in exp.bullets:
            content += f"  * {b}\n"
    content += f"\nSKILLS\n{', '.join(resume.sections.skills)}\n"
    return content
